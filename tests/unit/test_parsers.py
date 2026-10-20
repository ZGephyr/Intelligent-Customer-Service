"""文档解析器测试。"""

from __future__ import annotations

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from rag_service.core.exceptions import AppError, ErrorCode
from rag_service.infra.parsers import parse_file


def test_parse_file_utf8_txt_returns_single_block(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("七天无理由退货", encoding="utf-8")
    blocks = parse_file(p)
    assert len(blocks) == 1
    assert "七天无理由" in blocks[0].text


def test_parse_file_gbk_txt_falls_back_to_gbk(tmp_path):
    p = tmp_path / "b.txt"
    p.write_bytes("发票说明".encode("gbk"))
    assert "发票说明" in parse_file(p)[0].text


def test_parse_file_empty_markdown_raises_kb_empty_content(tmp_path):
    p = tmp_path / "empty.md"
    p.write_text("   ", encoding="utf-8")
    with pytest.raises(AppError) as ei:
        parse_file(p)
    assert ei.value.code == ErrorCode.KB_EMPTY_CONTENT


def test_parse_file_unsupported_suffix_raises(tmp_path):
    p = tmp_path / "a.png"
    p.write_bytes(b"x")
    with pytest.raises(AppError) as ei:
        parse_file(p)
    assert ei.value.code == ErrorCode.KB_FILE_TYPE_UNSUPPORTED


def test_parse_file_csv_renders_header_value_pairs(tmp_path):
    p = tmp_path / "p.csv"
    p.write_text("名称,价格\n耳机,199\n", encoding="utf-8")
    blocks = parse_file(p)
    assert "名称：耳机" in blocks[0].text
    assert "价格：199" in blocks[0].text


def test_parse_file_pdf_without_text_raises_kb_empty_content(tmp_path):
    path = tmp_path / "blank.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(path)
    doc.close()
    with pytest.raises(AppError) as ei:
        parse_file(path)
    assert ei.value.code == ErrorCode.KB_EMPTY_CONTENT


def test_parse_file_pdf_with_text_keeps_page_meta(tmp_path):
    path = tmp_path / "policy.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Warranty is 12 months")
    doc.save(path)
    doc.close()
    blocks = parse_file(path)
    assert "Warranty" in blocks[0].text
    assert blocks[0].meta["page"] == 1


def test_parse_file_docx_includes_table_cells(tmp_path):
    path = tmp_path / "p.docx"
    document = Document()
    document.add_paragraph("退换货政策")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "运费"
    table.rows[0].cells[1].text = "卖家承担"
    document.save(path)
    text = parse_file(path)[0].text
    assert "退换货政策" in text
    assert "卖家承担" in text


def test_parse_file_xlsx_renders_rows(tmp_path):
    path = tmp_path / "p.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["名称", "价格"])
    sheet.append(["耳机", 199])
    workbook.save(path)
    blocks = parse_file(path)
    assert "名称：耳机" in blocks[0].text
    assert "价格：199" in blocks[0].text


def test_parse_file_xlsx_header_only_raises_kb_empty_content(tmp_path):
    path = tmp_path / "only_header.xlsx"
    workbook = Workbook()
    workbook.active.append(["名称", "价格"])
    workbook.save(path)
    with pytest.raises(AppError) as ei:
        parse_file(path)
    assert ei.value.code == ErrorCode.KB_EMPTY_CONTENT
