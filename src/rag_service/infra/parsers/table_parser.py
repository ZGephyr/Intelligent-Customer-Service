"""表格解析：一行一个 block，渲染成「字段：值；」。"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook

from rag_service.core.exceptions import AppError, ErrorCode
from rag_service.core.types import ParsedBlock


def _render_row(header: Sequence[str], row: Iterable[Any]) -> str:
    pairs = []
    for name, value in zip(header, row):
        if value is None:
            continue
        text = str(value).strip()
        if text:
            pairs.append(f"{name.strip()}：{text}")
    return "；".join(pairs)


def _rows_to_blocks(rows: list[list[Any]], source_name: str) -> list[ParsedBlock]:
    if not rows:
        raise AppError(ErrorCode.KB_EMPTY_CONTENT, f"{source_name} 没有可解析的数据行")

    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    if not any(header):
        raise AppError(ErrorCode.KB_EMPTY_CONTENT, f"{source_name} 缺少表头")

    blocks: list[ParsedBlock] = []
    for index, row in enumerate(rows[1:], start=1):
        text = _render_row(header, row)
        if text:
            blocks.append(ParsedBlock(text=text, meta={"row": index}))

    if not blocks:
        raise AppError(ErrorCode.KB_EMPTY_CONTENT, f"{source_name} 只有表头，没有数据")
    return blocks


def parse_csv(path: Path) -> list[ParsedBlock]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            content = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise AppError(ErrorCode.KB_EMPTY_CONTENT, f"CSV 编码无法识别: {path.name}")

    rows = [row for row in csv.reader(io.StringIO(content)) if any(cell.strip() for cell in row)]
    return _rows_to_blocks(rows, path.name)


def parse_xlsx(path: Path) -> list[ParsedBlock]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise AppError(ErrorCode.KB_EMPTY_CONTENT, f"Excel 打开失败: {path.name}") from exc

    try:
        sheet = workbook.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    rows = [row for row in rows if any(cell is not None and str(cell).strip() for cell in row)]
    return _rows_to_blocks(rows, path.name)
