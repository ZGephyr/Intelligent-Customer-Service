"""Word 解析：段落 + 表格单元格。"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from rag_service.core.exceptions import AppError, ErrorCode
from rag_service.core.types import ParsedBlock


def parse(path: Path) -> list[ParsedBlock]:
    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise AppError(ErrorCode.KB_EMPTY_CONTENT, f"Word 打开失败: {path.name}") from exc

    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("；".join(cells))

    text = "\n".join(parts).strip()
    if not text:
        raise AppError(ErrorCode.KB_EMPTY_CONTENT, "Word 未解析出可用文本")
    return [ParsedBlock(text=text)]
