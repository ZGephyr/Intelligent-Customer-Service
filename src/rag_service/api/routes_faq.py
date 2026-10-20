"""管理端 FAQ 接口：列表、新增、编辑、删除、批量导入与统计。"""

from __future__ import annotations

import csv
import io
import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, File, Query, Request, UploadFile
from pydantic import BaseModel

from rag_service.api.deps import require_admin
from rag_service.api.schemas import FaqImportOut, FaqOut, FaqPageOut
from rag_service.core.exceptions import AppError, ErrorCode
from rag_service.core.logging import get_logger

logger = get_logger(__name__)

router = APIRouter(prefix="/api/admin/faqs", tags=["faq"], dependencies=[Depends(require_admin)])

COLUMN_MAP = {"问题": "question", "答案": "answer", "分类": "category", "别名": "aliases"}
SUPPORTED_IMPORT_SUFFIXES = {".csv", ".xlsx"}


class FaqCreateIn(BaseModel):
    question: str
    answer: str
    category: str | None = None
    aliases: str | None = None


class FaqUpdateIn(BaseModel):
    question: str | None = None
    answer: str | None = None
    category: str | None = None
    aliases: str | None = None
    enabled: bool | None = None


def _request_id(request: Request) -> str:
    return getattr(request.state, "request_id", uuid.uuid4().hex)


def _parse_import_rows(filename: str, data: bytes) -> list[dict[str, str]]:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_IMPORT_SUFFIXES:
        raise AppError(
            ErrorCode.FAQ_IMPORT_INVALID,
            "仅支持 .csv 或 .xlsx 导入",
            http_status=400,
        )
    if not data:
        raise AppError(ErrorCode.FAQ_IMPORT_INVALID, "导入文件为空", http_status=400)

    csv.field_size_limit(1_000_000)
    try:
        raw_rows = _read_rows(suffix, data)
    except AppError:
        raise
    except Exception as exc:  # noqa: BLE001 - 损坏文件统一转 400
        logger.warning("FAQ 导入文件解析失败: file=%s error=%s", filename, exc)
        raise AppError(
            ErrorCode.FAQ_IMPORT_INVALID,
            "文件解析失败，请确认是有效的 CSV 或 XLSX",
            http_status=400,
        ) from exc

    rows: list[dict[str, str]] = []
    for raw in raw_rows:
        normalized: dict[str, str] = {}
        for key, value in raw.items():
            if key is None:
                continue
            column = COLUMN_MAP.get(str(key).strip(), str(key).strip())
            normalized[column] = "" if value is None else str(value).strip()
        if normalized.get("question") or normalized.get("answer"):
            rows.append(normalized)

    if not rows:
        raise AppError(
            ErrorCode.FAQ_IMPORT_INVALID,
            "没有可用数据：表头需包含「问题」和「答案」",
            http_status=400,
        )
    if not any("question" in row for row in rows):
        raise AppError(
            ErrorCode.FAQ_IMPORT_INVALID,
            "表头需包含「问题」和「答案」列",
            http_status=400,
        )
    return rows


def _read_rows(suffix: str, data: bytes) -> list[dict[str, str]]:
    if suffix == ".csv":
        text = None
        for encoding in ("utf-8-sig", "gbk"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise AppError(ErrorCode.FAQ_IMPORT_INVALID, "CSV 编码无法识别，请另存为 UTF-8", 400)
        reader = csv.DictReader(io.StringIO(text))
        return [dict(row) for row in reader]

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(iterator, [])]
        return [
            dict(zip(header, ["" if cell is None else str(cell) for cell in row], strict=False))
            for row in iterator
        ]
    finally:
        workbook.close()


@router.get("")
def list_faqs(
    request: Request,
    page: int = Query(default=1, ge=1),
    size: int = Query(default=20, ge=1, le=100),
    keyword: str = "",
    category: str = "",
    enabled: bool | None = None,
) -> dict:
    service = request.app.state.services.faq_admin
    items, total = service.list(page=page, size=size, keyword=keyword, category=category, enabled=enabled)
    payload = FaqPageOut(
        items=[FaqOut.model_validate(item) for item in items],
        total=total,
        page=page,
        size=size,
    )
    return {"code": 0, "data": payload.model_dump(mode="json"), "request_id": _request_id(request)}


@router.post("")
def create_faq(request: Request, body: FaqCreateIn) -> dict:
    service = request.app.state.services.faq_admin
    row = service.create(body.question, body.answer, category=body.category, aliases=body.aliases)
    return {
        "code": 0,
        "data": FaqOut.model_validate(row).model_dump(mode="json"),
        "request_id": _request_id(request),
    }


@router.put("/{faq_id}")
def update_faq(request: Request, faq_id: str, body: FaqUpdateIn) -> dict:
    service = request.app.state.services.faq_admin
    row = service.update(
        faq_id,
        question=body.question,
        answer=body.answer,
        category=body.category,
        aliases=body.aliases,
        enabled=body.enabled,
    )
    return {
        "code": 0,
        "data": FaqOut.model_validate(row).model_dump(mode="json"),
        "request_id": _request_id(request),
    }


@router.delete("/{faq_id}")
def delete_faq(request: Request, faq_id: str) -> dict:
    request.app.state.services.faq_admin.delete(faq_id)
    return {"code": 0, "data": {"deleted": faq_id}, "request_id": _request_id(request)}


@router.post("/import")
async def import_faqs(request: Request, file: UploadFile = File(...)) -> dict:
    data = await file.read()
    rows = _parse_import_rows(file.filename or "", data)
    result = request.app.state.services.faq_admin.import_rows(rows)
    payload = FaqImportOut(created=result.created, updated=result.updated, skipped=result.skipped)
    logger.info("FAQ 导入: 文件=%s 行数=%s", file.filename, len(rows))
    return {"code": 0, "data": payload.model_dump(), "request_id": _request_id(request)}


@router.get("/stats")
def faq_stats(request: Request) -> dict:
    return {
        "code": 0,
        "data": request.app.state.services.faq_admin.stats(),
        "request_id": _request_id(request),
    }
